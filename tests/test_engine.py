"""Engine contract: genesis hash, second-row link, UNOWNED, RENAMED, tamper, missing file, import hash."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from employeelock.engine import (
    GENESIS_PREV,
    LOG_DATA_START,
    append_row,
    guess_kind,
    import_files,
    init_workbook,
    row_hash,
    sha256_file,
    verify_workbook,
)


def test_guess_kind() -> None:
    assert guess_kind(Path("a.wav")) == "audio"
    assert guess_kind(Path("a.mp4")) == "video"
    assert guess_kind(Path("a.png")) == "photo"
    assert guess_kind(Path("a.md")) == "document"
    assert guess_kind(Path("a.bin")) == "other"
    assert guess_kind(Path("a.bin"), "audio") == "audio"


def test_genesis_hash(tmp_path: Path) -> None:
    path = tmp_path / "el.xlsx"
    init_workbook(path, demo=False)
    rec = append_row(
        path,
        event="genesis event",
        result="genesis result",
        blame="",
        owner="",
        short="soon",
        long="lasted",
        timestamp="2026-09-02T15:00:00Z",
        confidence=0.7,
    )
    assert rec["prev_hash"] == GENESIS_PREV
    assert rec["entry_id"] == "EL-0001"
    _, digest = row_hash(
        {
            "entry_id": "EL-0001",
            "timestamp": "2026-09-02T15:00:00Z",
            "event": "genesis event",
            "result": "genesis result",
            "blame_placed": "",
            "owner_named": "",
            "renamed_from": "",
            "outcome_short": "soon",
            "outcome_long": "lasted",
            "evidence_ids": "",
            "confidence": 0.7,
            "observer": "operator",
            "file_sha256s": "",
            "prev_hash": GENESIS_PREV,
        }
    )
    assert rec["row_hash"] == digest
    assert rec["unowned"] == "UNOWNED"
    report = verify_workbook(path)
    assert report["ok"] is True
    assert report["rows"] == 1
    assert report["unowned"] == 1


def test_second_row_link(tmp_path: Path) -> None:
    path = tmp_path / "el.xlsx"
    init_workbook(path, demo=False)
    first = append_row(
        path,
        event="one",
        result="r1",
        timestamp="2026-09-02T15:00:00Z",
    )
    second = append_row(
        path,
        event="two",
        result="r2",
        owner="records desk",
        timestamp="2026-09-02T15:00:01Z",
    )
    assert second["prev_hash"] == first["row_hash"]
    assert second["entry_id"] == "EL-0002"
    report = verify_workbook(path)
    assert report["ok"] is True
    assert report["rows"] == 2


def test_blank_owner_unowned(tmp_path: Path) -> None:
    path = tmp_path / "el.xlsx"
    init_workbook(path, demo=False)
    rec = append_row(path, event="e", result="r", owner="")
    assert rec["unowned"] == "UNOWNED"
    wb = load_workbook(path)
    formula = str(wb["LOG"].cell(LOG_DATA_START, 7).value)
    assert "UNOWNED" in formula
    assert "TRIM(F5)" in formula
    report = verify_workbook(path)
    assert report["unowned"] == 1


def test_renamed_from_flag(tmp_path: Path) -> None:
    path = tmp_path / "el.xlsx"
    init_workbook(path, demo=False)
    rec = append_row(
        path,
        event="e",
        result="r",
        owner="records desk",
        renamed_from="ticket queue",
    )
    assert rec["name_moved"] == "RENAMED"
    wb = load_workbook(path)
    formula = str(wb["LOG"].cell(LOG_DATA_START, 9).value)
    assert "RENAMED" in formula
    assert "TRIM(H5)" in formula


def test_verify_catches_edited_event(tmp_path: Path) -> None:
    path = tmp_path / "el.xlsx"
    init_workbook(path, demo=False)
    append_row(path, event="keep", result="r", timestamp="2026-09-02T15:00:00Z")
    wb = load_workbook(path)
    wb["LOG"].cell(LOG_DATA_START, 3, "edited event")
    wb.save(path)
    report = verify_workbook(path)
    assert report["ok"] is False
    assert report["errors"]
    assert "row_hash" in report["errors"][0]


def test_missing_file_reported_ok_stays(tmp_path: Path) -> None:
    path = tmp_path / "el.xlsx"
    blob = tmp_path / "note.txt"
    blob.write_text("bytes", encoding="utf-8")
    init_workbook(path, demo=False)
    import_files(
        path,
        [blob],
        event="imported",
        result="hashed",
        owner="records desk",
        timestamp="2026-09-02T15:00:00Z",
    )
    blob.unlink()
    report = verify_workbook(path)
    assert report["ok"] is True
    assert report["missing_files"]
    assert report["rows"] == 1


def test_import_hashes_file_bytes(tmp_path: Path) -> None:
    path = tmp_path / "el.xlsx"
    blob = tmp_path / "spec.md"
    blob.write_text("# EmployeeLock spec\n", encoding="utf-8")
    expected, size = sha256_file(blob)
    init_workbook(path, demo=False)
    rec = import_files(
        path,
        [blob],
        event="import spec",
        result="indexed",
        owner="records desk",
        timestamp="2026-09-02T15:00:00Z",
    )
    assert rec["evidence"][0]["file_sha256"] == expected
    assert rec["evidence"][0]["size_bytes"] == size
    assert expected in rec["canonical"] or expected in rec["row_hash"] or rec["evidence"]
    wb = load_workbook(path)
    stored = str(wb["EVIDENCE"].cell(5, 7).value)
    assert stored == expected
    assert str(wb["LOG"].cell(LOG_DATA_START, 16).value) == "EV-0001"
    assert str(wb["LOG"].cell(LOG_DATA_START, 17).value) == expected


def test_demo_init_two_rows(tmp_path: Path) -> None:
    path = tmp_path / "el.xlsx"
    init_workbook(path, demo=True)
    report = verify_workbook(path)
    assert report["ok"] is True
    assert report["rows"] == 2
    assert report["unowned"] == 1
    wb = load_workbook(path)
    assert "COVER" in wb.sheetnames
    assert "LOG" in wb.sheetnames
    assert "EVIDENCE" in wb.sheetnames
    assert "CHAIN" in wb.sheetnames
    assert "OWNERS" in wb.sheetnames
    assert "DASH" in wb.sheetnames
    assert "LISTS" in wb.sheetnames
    event1 = str(wb["LOG"].cell(LOG_DATA_START, 3).value)
    event2 = str(wb["LOG"].cell(LOG_DATA_START + 1, 3).value)
    assert "unnamed" in event1.lower() or "no named owner" in event1.lower()
    assert "records desk" in event2.lower() or "records desk" in str(wb["LOG"].cell(LOG_DATA_START + 1, 6).value).lower()


def test_confidence_six_decimals() -> None:
    canon, _digest = row_hash(
        {
            "entry_id": "EL-0001",
            "timestamp": "2026-09-02T15:00:00Z",
            "event": "e",
            "result": "r",
            "blame_placed": "",
            "owner_named": "",
            "renamed_from": "",
            "outcome_short": "",
            "outcome_long": "",
            "evidence_ids": "",
            "confidence": 0.7,
            "observer": "operator",
            "file_sha256s": "",
            "prev_hash": GENESIS_PREV,
        }
    )
    assert '"confidence":0.700000' in canon
    assert canon == (
        '{"blame_placed":"","confidence":0.700000,"entry_id":"EL-0001",'
        '"event":"e","evidence_ids":"","file_sha256s":"","observer":"operator",'
        '"outcome_long":"","outcome_short":"","owner_named":"",'
        f'"prev_hash":"{GENESIS_PREV}","renamed_from":"","result":"r",'
        '"timestamp":"2026-09-02T15:00:00Z"}'
    )


def test_no_copy_to_and_log_columns(tmp_path: Path) -> None:
    path = tmp_path / "el.xlsx"
    init_workbook(path, demo=False)
    wb = load_workbook(path)
    headers = [wb["LOG"].cell(4, c).value for c in range(1, 22)]
    assert headers[0] == "entry_id"
    assert headers[20] == "link_ok"
    assert headers[6] == "unowned"
    cover = "\n".join(str(c.value or "") for row in wb["COVER"].iter_rows() for c in row)
    assert "no --copy-to" in cover.lower() or "no --copy-to" in cover
