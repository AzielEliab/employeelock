"""Self-check for EmployeeLock. NASA-robust, no network, no telemetry.

    employeelock doctor
"""

from __future__ import annotations

import json
import os
import tempfile
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook

from employeelock import __version__
from employeelock.engine import (
    ENGINE_VERSION,
    GENESIS_PREV,
    LIMITATION,
    LOG_DATA_START,
    MAX_LOG,
    SPEC_STRING,
    append_row,
    import_files,
    init_workbook,
    row_hash,
    verify_workbook,
)
from employeelock.ui import LOOPBACK, make_server

Check = tuple[str, bool, str]


def _ok(name: str, detail: str = "") -> Check:
    return name, True, detail


def _fail(name: str, detail: str) -> Check:
    return name, False, detail


def _check_version() -> Check:
    if __version__ == ENGINE_VERSION == "0.1.0":
        return _ok("version", __version__)
    return _fail("version", f"{__version__} vs engine {ENGINE_VERSION}")


def _check_spec() -> Check:
    if SPEC_STRING == "employeelock-v0":
        return _ok("spec", SPEC_STRING)
    return _fail("spec", SPEC_STRING)


def _check_genesis() -> Check:
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "el.xlsx"
        init_workbook(path, demo=False)
        rec = append_row(
            path,
            event="genesis event",
            result="genesis result",
            blame="",
            owner="",
            short="soon",
            long="lasted",
            timestamp="2026-09-02T15:00:00Z",
        )
        if rec["prev_hash"] != GENESIS_PREV:
            return _fail("genesis prev_hash", rec["prev_hash"])
        expected = row_hash(
            {
                "entry_id": rec["entry_id"],
                "timestamp": "2026-09-02T15:00:00Z",
                "event": "genesis event",
                "result": "genesis result",
                "blame_placed": "",
                "owner_named": "",
                "renamed_from": "",
                "outcome_short": "soon",
                "outcome_long": "lasted",
                "evidence_ids": "",
                "confidence": 0.7,
                "observer": "operator",
                "file_sha256s": "",
                "prev_hash": GENESIS_PREV,
            }
        )[1]
        if rec["row_hash"] != expected:
            return _fail("genesis hash", f"{rec['row_hash']} != {expected}")
        if rec["unowned"] != "UNOWNED":
            return _fail("genesis unowned", rec["unowned"])
        return _ok("genesis", rec["row_hash"][:16] + "…")


def _check_link_and_flags() -> Check:
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "el.xlsx"
        init_workbook(path, demo=False)
        first = append_row(
            path,
            event="one",
            result="r1",
            owner="",
            timestamp="2026-09-02T15:00:00Z",
        )
        second = append_row(
            path,
            event="two",
            result="r2",
            owner="records desk",
            renamed_from="ticket queue",
            timestamp="2026-09-02T15:00:01Z",
        )
        if second["prev_hash"] != first["row_hash"]:
            return _fail("second-row link", f"{second['prev_hash']} != {first['row_hash']}")
        if second["name_moved"] != "RENAMED":
            return _fail("renamed flag", second["name_moved"])
        wb = load_workbook(path)
        g = str(wb["LOG"].cell(LOG_DATA_START, 7).value)
        i = str(wb["LOG"].cell(LOG_DATA_START + 1, 9).value)
        if "UNOWNED" not in g:
            return _fail("unowned formula", g)
        if "RENAMED" not in i:
            return _fail("renamed formula", i)
        return _ok("link+flags", "second prev_hash matches; UNOWNED/RENAMED formulas")


def _check_tamper() -> Check:
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "el.xlsx"
        init_workbook(path, demo=False)
        append_row(path, event="keep", result="r", timestamp="2026-09-02T15:00:00Z")
        wb = load_workbook(path)
        wb["LOG"].cell(LOG_DATA_START, 3, "edited event")
        wb.save(path)
        report = verify_workbook(path)
        if report["ok"]:
            return _fail("tamper", "verify did not catch edited event")
        if not report["errors"]:
            return _fail("tamper", "no errors listed")
        return _ok("tamper", report["errors"][0][:80])


def _check_import_and_missing() -> Check:
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "el.xlsx"
        blob = Path(tmp) / "note.txt"
        blob.write_text("employee lock format proof\n", encoding="utf-8")
        init_workbook(path, demo=False)
        rec = import_files(
            path,
            [blob],
            event="imported file",
            result="hashed",
            owner="records desk",
            timestamp="2026-09-02T15:00:00Z",
        )
        if not rec.get("evidence"):
            return _fail("import", "no evidence rows")
        digest = rec["evidence"][0]["file_sha256"]
        if len(digest) != 64:
            return _fail("import hash", digest)
        blob.unlink()
        report = verify_workbook(path)
        if not report["ok"]:
            return _fail("missing file failed ok", str(report["errors"]))
        if not report["missing_files"]:
            return _fail("missing file", "not reported")
        return _ok("import+missing", f"sha256 {digest[:16]}…; missing reported; ok={report['ok']}")


def _check_loopback() -> Check:
    try:
        make_server("0.0.0.0", 9)
    except ValueError as exc:
        if "loopback" in str(exc).lower() and "127.0.0.1" in LOOPBACK:
            return _ok("loopback", "rejects 0.0.0.0")
        return _fail("loopback", str(exc))
    return _fail("loopback", "accepted 0.0.0.0")


def _check_capacity_formulas() -> Check:
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "el.xlsx"
        init_workbook(path, demo=False)
        wb = load_workbook(path)
        last = LOG_DATA_START + MAX_LOG - 1
        u_last = str(wb["LOG"].cell(last, 21).value)
        if f"T{last - 1}" not in u_last:
            return _fail("MAX_LOG formulas", u_last)
        return _ok("MAX_LOG", str(MAX_LOG))


CHECKS: tuple[Callable[[], Check], ...] = (
    _check_version,
    _check_spec,
    _check_genesis,
    _check_link_and_flags,
    _check_tamper,
    _check_import_and_missing,
    _check_loopback,
    _check_capacity_formulas,
)


def run_doctor(*, as_json: bool = False) -> int:
    results = []
    failed = 0
    for fn in CHECKS:
        name, ok, detail = fn()
        results.append({"name": name, "ok": ok, "detail": detail})
        if not ok:
            failed += 1
        mark = "ok" if ok else "FAIL"
        if not as_json:
            print(f"[{mark}] {name}" + (f" — {detail}" if detail else ""))
    payload = {
        "ok": failed == 0,
        "failed": failed,
        "checks": results,
        "version": __version__,
        "spec": SPEC_STRING,
        "limitation": LIMITATION,
        "network": False,
        "telemetry": False,
    }
    if as_json:
        print(json.dumps(payload, indent=2))
    else:
        print("limitation:", LIMITATION)
        print("doctor", "passed" if failed == 0 else "failed")
    return 0 if failed == 0 else 1
