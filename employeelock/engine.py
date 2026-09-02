"""EmployeeLock v0 engine. Path B: CLI appends, never rewrites hashed cells.

Spec string: employeelock-v0. Paper: EL-WP-0.1. Author: Aziel Eliab.
"""

from __future__ import annotations

import hashlib
import json
import os
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

ENGINE_VERSION = "0.1.0"
SPEC_STRING = "employeelock-v0"
PAPER_ID = "EL-WP-0.1"
MAX_LOG = 500
MAX_EVIDENCE = 500
GENESIS_PREV = "0" * 64
LOG_HEADER_ROW = 4
LOG_DATA_START = 5
EVIDENCE_HEADER_ROW = 4
EVIDENCE_DATA_START = 5
OWNERS_DATA_START = 5
CONF_PLACEHOLDER = "__EL_CONFIDENCE__"
DEFAULT_CONFIDENCE = 0.7
DEFAULT_OBSERVER = "operator"
DEFAULT_EXHIBIT = "N"

LIMITATION = (
    "THIS IS: workbook (COVER, LOG, EVIDENCE, CHAIN, OWNERS, DASH, LISTS) "
    "+ CLI (init/append/import/verify) + linear hash chain + countermeasure "
    "against unowned/renamed rows. "
    "THIS IS NOT: UL or a BAL issue paper; FoldLock; TemporalLock "
    "(borrows ethic, different product); court filing / exhibit stickerer / "
    "counsel; truth score / consensus / token; remote uploader / anonymous "
    "relay; a charge sheet against a named living person. "
    "Demo rows are generic format proof, not case facts. "
    "Hosted API never stores xlsx."
)

HASHED_FIELDS: tuple[str, ...] = (
    "entry_id",
    "timestamp",
    "event",
    "result",
    "blame_placed",
    "owner_named",
    "renamed_from",
    "outcome_short",
    "outcome_long",
    "evidence_ids",
    "confidence",
    "observer",
    "file_sha256s",
    "prev_hash",
)

LOG_HEADERS: tuple[str, ...] = (
    "entry_id",
    "timestamp_utc",
    "event",
    "result",
    "blame_placed",
    "owner_named",
    "unowned",
    "renamed_from",
    "name_moved",
    "outcome_short",
    "outcome_long",
    "cites",
    "confidence",
    "observer",
    "exhibit_potential",
    "evidence_ids",
    "file_sha256s",
    "canonical",
    "prev_hash",
    "row_hash",
    "link_ok",
)

EVIDENCE_HEADERS: tuple[str, ...] = (
    "evidence_id",
    "imported_utc",
    "kind",
    "original_name",
    "stored_path",
    "size_bytes",
    "file_sha256",
    "linked_entry_id",
    "note",
)

CHAIN_HEADERS: tuple[str, ...] = (
    "entry_id",
    "timestamp_utc",
    "prev_hash",
    "row_hash",
    "link_ok",
)

KIND_SUFFIXES: dict[str, tuple[str, ...]] = {
    "audio": (".wav", ".mp3", ".m4a", ".flac", ".ogg", ".aac"),
    "video": (".mp4", ".mov", ".mkv", ".webm", ".avi"),
    "photo": (".jpg", ".jpeg", ".png", ".gif", ".webp", ".tif", ".tiff", ".heic"),
    "document": (".pdf", ".doc", ".docx", ".txt", ".md", ".xlsx", ".csv"),
}

# Colors (paper: UNOWNED/BREAK red, owned/OK green, RENAMED amber, GENESIS gold-tan)
FILL_RED = PatternFill("solid", fgColor="C0392B")
FILL_GREEN = PatternFill("solid", fgColor="1E8449")
FILL_AMBER = PatternFill("solid", fgColor="D68910")
FILL_GOLD = PatternFill("solid", fgColor="C9A227")
FONT_WHITE = Font(color="FFFFFF", name="Calibri", size=11)
FONT_BLUE = Font(name="Calibri", size=11, color="0000CC")
FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="E8E0D0")
FONT_BODY = Font(name="Calibri", size=11, color="1A1A1A")
FONT_FORMULA = Font(name="Calibri", size=11, color="1A1A1A")
FILL_HEADER = PatternFill("solid", fgColor="1A1610")
FILL_COVER = PatternFill("solid", fgColor="0B0B0B")
THIN = Border(
    left=Side(style="thin", color="2A261C"),
    right=Side(style="thin", color="2A261C"),
    top=Side(style="thin", color="2A261C"),
    bottom=Side(style="thin", color="2A261C"),
)
INPUT_COLS = {3, 4, 5, 6, 8, 10, 11, 12, 13, 14, 15}  # C-F, H, J-O (1-indexed)

DEMO_ROW_1 = {
    "event": "process outcome recorded with no named owner",
    "result": "row logged as format proof",
    "blame": "",
    "owner": "",
    "short": "unnamed prior process visible as UNOWNED",
    "long": "demo genesis row; replace with real work. not an accusation.",
    "renamed_from": "",
    "confidence": 0.700000,
    "observer": "operator",
    "exhibit": "N",
}
DEMO_ROW_2 = {
    "event": "records desk took the ticket from the prior queue name",
    "result": "owner named; prior name moved",
    "blame": "",
    "owner": "records desk",
    "short": "RENAMED flag set; owned as a record",
    "long": "demo owned row; replace with real work. not an accusation.",
    "renamed_from": "ticket queue",
    "confidence": 0.700000,
    "observer": "operator",
    "exhibit": "N",
}


def utc_now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _as_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def format_confidence(confidence: Any) -> str:
    return f"{float(confidence):.6f}"


def canonical_json(fields: dict[str, Any]) -> str:
    """UTF-8 JSON, sorted keys, compact separators. confidence is a 6-decimal number."""
    payload: dict[str, Any] = {}
    for key in HASHED_FIELDS:
        if key == "confidence":
            payload[key] = CONF_PLACEHOLDER
        else:
            payload[key] = _as_str(fields.get(key, ""))
    raw = json.dumps(payload, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return raw.replace(f'"{CONF_PLACEHOLDER}"', format_confidence(fields.get("confidence", 0.0)))


def row_hash(fields: dict[str, Any]) -> tuple[str, str]:
    """Return (canonical, sha256 hex)."""
    canon = canonical_json(fields)
    digest = hashlib.sha256(canon.encode("utf-8")).hexdigest()
    return canon, digest


def sha256_file(path: Path) -> tuple[str, int]:
    data = path.read_bytes()
    return hashlib.sha256(data).hexdigest(), len(data)


def guess_kind(path: Path, override: str | None = None) -> str:
    if override:
        return override.strip().lower()
    suffix = path.suffix.lower()
    for kind, suffixes in KIND_SUFFIXES.items():
        if suffix in suffixes:
            return kind
    return "other"


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip() if not isinstance(value, str) else str(value)


def _log_formula_unowned(row: int) -> str:
    return f'=IF(C{row}="","",IF(TRIM(F{row})="","UNOWNED","owned"))'


def _log_formula_renamed(row: int) -> str:
    return f'=IF(C{row}="","",IF(TRIM(H{row})="","stable","RENAMED"))'


def _log_formula_link(row: int) -> str:
    if row == LOG_DATA_START:
        return f'=IF(C{row}="","",IF(S{row}="{GENESIS_PREV}","GENESIS","BROKEN"))'
    return f'=IF(C{row}="","",IF(S{row}=T{row - 1},"OK","BREAK"))'


def _style_header(ws: Worksheet, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row, col)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        cell.border = THIN


def _apply_log_cf(ws: Worksheet) -> None:
    last = LOG_DATA_START + MAX_LOG - 1
    g_range = f"G{LOG_DATA_START}:G{last}"
    i_range = f"I{LOG_DATA_START}:I{last}"
    u_range = f"U{LOG_DATA_START}:U{last}"
    ws.conditional_formatting.add(g_range, CellIsRule(operator="equal", formula=['"UNOWNED"'], fill=FILL_RED, font=FONT_WHITE))
    ws.conditional_formatting.add(g_range, CellIsRule(operator="equal", formula=['"owned"'], fill=FILL_GREEN, font=FONT_WHITE))
    ws.conditional_formatting.add(i_range, CellIsRule(operator="equal", formula=['"RENAMED"'], fill=FILL_AMBER, font=FONT_WHITE))
    ws.conditional_formatting.add(u_range, CellIsRule(operator="equal", formula=['"BREAK"'], fill=FILL_RED, font=FONT_WHITE))
    ws.conditional_formatting.add(u_range, CellIsRule(operator="equal", formula=['"BROKEN"'], fill=FILL_RED, font=FONT_WHITE))
    ws.conditional_formatting.add(u_range, CellIsRule(operator="equal", formula=['"OK"'], fill=FILL_GREEN, font=FONT_WHITE))
    ws.conditional_formatting.add(u_range, CellIsRule(operator="equal", formula=['"GENESIS"'], fill=FILL_GOLD, font=FONT_WHITE))


def _print_tabloid_landscape(ws: Worksheet) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _build_cover(ws: Worksheet) -> None:
    ws.sheet_view.showGridLines = False
    ws["A1"] = "EmployeeLock"
    ws["A1"].font = Font(name="Calibri", size=22, bold=True, color="C9A227")
    ws["A2"] = "A hash-chained accountability workbook"
    ws["A2"].font = Font(name="Calibri", size=14, italic=True, color="8A8070")
    lines = [
        (4, "Spec"),
        (5, f"Spec string: {SPEC_STRING}"),
        (6, f"Paper: {PAPER_ID}  ·  Product: EmployeeLock v0  ·  License: Apache-2.0"),
        (7, "Author: Aziel Eliab  ·  2 September 2026"),
        (9, "THIS IS"),
        (10, "workbook (COVER, LOG, EVIDENCE, CHAIN, OWNERS, DASH, LISTS)"),
        (11, "CLI (init / append / import / verify)"),
        (12, "a linear hash chain of rows"),
        (13, "a countermeasure against unowned / renamed rows (not a member of UL-CAT)"),
        (15, "THIS IS NOT"),
        (16, "UL, and not a BAL issue paper"),
        (17, "FoldLock (compression)"),
        (18, "TemporalLock (borrows the ethic; different product)"),
        (19, "a court filing, exhibit stickerer, or counsel"),
        (20, "a truth score, consensus layer, or token"),
        (21, "a remote uploader or anonymous relay"),
        (22, "a charge sheet against a named living person in the shipped demo"),
        (24, "Operator"),
        (25, "Name the event. Name who owns the record. Keep leftover blame blank if it is blank. Chain the row."),
        (26, "Blame and ownership are not the same field."),
        (27, "Media is hashed, not embedded. The workbook is an index. v0 has no --copy-to."),
        (28, "Append, don't edit hashed cells. A later long outcome is a new row that cites the old row_hash."),
        (29, "Verify bytes, not narratives. verify recomputes hashes. It does not decide who was right."),
        (31, "CLI"),
        (32, "python3 employeelock.py init WORKBOOK.xlsx"),
        (33, "python3 employeelock.py append WORKBOOK.xlsx --event TEXT --result TEXT --blame TEXT --owner TEXT --short TEXT --long TEXT"),
        (34, "python3 employeelock.py import WORKBOOK.xlsx FILE [FILE ...] --event TEXT ..."),
        (35, "python3 employeelock.py verify WORKBOOK.xlsx"),
        (36, "employeelock ui     # http://127.0.0.1:8871 loopback only"),
        (37, "employeelock doctor"),
        (39, "Demo rows after init are generic format proof (records desk / unnamed prior process). Replace them. Not accusations."),
        (40, "Do not put a legal name or home location in this template. Public identity: Aziel Eliab."),
        (41, "Forks are welcome and always allowed."),
    ]
    for row, text in lines:
        ws.cell(row, 1, text)
        if text in {"Spec", "THIS IS", "THIS IS NOT", "Operator", "CLI"}:
            ws.cell(row, 1).font = Font(name="Calibri", size=12, bold=True, color="C9A227")
        else:
            ws.cell(row, 1).font = Font(name="Calibri", size=11, color="E8E0D0")
    ws.column_dimensions["A"].width = 110
    ws.row_dimensions[1].height = 28
    _print_tabloid_landscape(ws)
    ws.sheet_properties.tabColor = "C9A227"


def _build_log(ws: Worksheet) -> None:
    ws["A1"] = "LOG — append-only event register"
    ws["A2"] = f"{SPEC_STRING}  ·  MAX_LOG={MAX_LOG}  ·  headers at row 4  ·  data from row 5"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="C9A227")
    ws["A2"].font = Font(name="Calibri", size=10, color="8A8070")
    for col, name in enumerate(LOG_HEADERS, start=1):
        ws.cell(LOG_HEADER_ROW, col, name)
    _style_header(ws, LOG_HEADER_ROW, len(LOG_HEADERS))
    last = LOG_DATA_START + MAX_LOG - 1
    for row in range(LOG_DATA_START, last + 1):
        ws.cell(row, 7, _log_formula_unowned(row)).font = FONT_FORMULA
        ws.cell(row, 9, _log_formula_renamed(row)).font = FONT_FORMULA
        ws.cell(row, 21, _log_formula_link(row)).font = FONT_FORMULA
        for col in range(1, 22):
            cell = ws.cell(row, col)
            cell.border = THIN
            if col in INPUT_COLS:
                cell.font = FONT_BLUE
    widths = {
        1: 12, 2: 22, 3: 36, 4: 28, 5: 18, 6: 18, 7: 12, 8: 18, 9: 12,
        10: 24, 11: 28, 12: 16, 13: 12, 14: 12, 15: 12, 16: 16, 17: 20,
        18: 40, 19: 20, 20: 20, 21: 12,
    }
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A{LOG_HEADER_ROW}:U{last}"
    _apply_log_cf(ws)
    _print_tabloid_landscape(ws)
    observer_dv = DataValidation(type="list", formula1="=LISTS!$A$2:$A$5", allow_blank=True)
    exhibit_dv = DataValidation(type="list", formula1="=LISTS!$B$2:$B$4", allow_blank=True)
    observer_dv.add(f"N{LOG_DATA_START}:N{last}")
    exhibit_dv.add(f"O{LOG_DATA_START}:O{last}")
    ws.add_data_validation(observer_dv)
    ws.add_data_validation(exhibit_dv)


def _build_evidence(ws: Worksheet) -> None:
    ws["A1"] = "EVIDENCE — file index (hashed, not embedded)"
    ws["A2"] = f"{SPEC_STRING}  ·  MAX_EVIDENCE={MAX_EVIDENCE}  ·  operator keeps the bytes  ·  no --copy-to"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="C9A227")
    ws["A2"].font = Font(name="Calibri", size=10, color="8A8070")
    for col, name in enumerate(EVIDENCE_HEADERS, start=1):
        ws.cell(EVIDENCE_HEADER_ROW, col, name)
    _style_header(ws, EVIDENCE_HEADER_ROW, len(EVIDENCE_HEADERS))
    last = EVIDENCE_DATA_START + MAX_EVIDENCE - 1
    for col, width in enumerate((14, 22, 12, 28, 48, 14, 66, 16, 28), start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    kind_dv = DataValidation(type="list", formula1="=LISTS!$C$2:$C$7", allow_blank=True)
    kind_dv.add(f"C{EVIDENCE_DATA_START}:C{last}")
    ws.add_data_validation(kind_dv)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A{EVIDENCE_HEADER_ROW}:I{last}"
    _print_tabloid_landscape(ws)


def _build_chain(ws: Worksheet) -> None:
    ws["A1"] = "CHAIN — link fields without the working columns"
    ws["A2"] = "link_ok pulled from LOG"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="C9A227")
    for col, name in enumerate(CHAIN_HEADERS, start=1):
        ws.cell(LOG_HEADER_ROW, col, name)
    _style_header(ws, LOG_HEADER_ROW, len(CHAIN_HEADERS))
    last = LOG_DATA_START + MAX_LOG - 1
    for row in range(LOG_DATA_START, last + 1):
        ws.cell(row, 1, f"=LOG!A{row}")
        ws.cell(row, 2, f"=LOG!B{row}")
        ws.cell(row, 3, f"=LOG!S{row}")
        ws.cell(row, 4, f"=LOG!T{row}")
        ws.cell(row, 5, f"=LOG!U{row}")
    for col, width in enumerate((14, 22, 66, 66, 12), start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    u_range = f"E{LOG_DATA_START}:E{last}"
    ws.conditional_formatting.add(u_range, CellIsRule(operator="equal", formula=['"BREAK"'], fill=FILL_RED, font=FONT_WHITE))
    ws.conditional_formatting.add(u_range, CellIsRule(operator="equal", formula=['"BROKEN"'], fill=FILL_RED, font=FONT_WHITE))
    ws.conditional_formatting.add(u_range, CellIsRule(operator="equal", formula=['"OK"'], fill=FILL_GREEN, font=FONT_WHITE))
    ws.conditional_formatting.add(u_range, CellIsRule(operator="equal", formula=['"GENESIS"'], fill=FILL_GOLD, font=FONT_WHITE))
    ws.freeze_panes = "A5"
    _print_tabloid_landscape(ws)


def _build_owners(ws: Worksheet) -> None:
    ws["A1"] = "OWNERS — distinct owner_named the CLI has seen"
    ws["A2"] = "Blank owner_named is UNOWNED. renamed_from filled is RENAMED. Not a charge sheet."
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="C9A227")
    headers = ("owner_named", "rows_owned", "blame_mentions")
    for col, name in enumerate(headers, start=1):
        ws.cell(4, col, name)
    _style_header(ws, 4, 3)
    ws["E4"] = "surface"
    ws["F4"] = "count"
    _style_header(ws, 4, 6)
    ws["E5"] = "UNOWNED"
    ws["F5"] = f'=COUNTIF(LOG!G{LOG_DATA_START}:G{LOG_DATA_START + MAX_LOG - 1},"UNOWNED")'
    ws["E6"] = "RENAMED"
    ws["F6"] = f'=COUNTIF(LOG!I{LOG_DATA_START}:I{LOG_DATA_START + MAX_LOG - 1},"RENAMED")'
    ws["E5"].fill = FILL_RED
    ws["E5"].font = FONT_WHITE
    ws["E6"].fill = FILL_AMBER
    ws["E6"].font = FONT_WHITE
    for col, width in enumerate((24, 14, 16, 8, 14, 12), start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    _print_tabloid_landscape(ws)


def _build_dash(ws: Worksheet) -> None:
    ws["A1"] = "DASH — counts"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="C9A227")
    last = LOG_DATA_START + MAX_LOG - 1
    ev_last = EVIDENCE_DATA_START + MAX_EVIDENCE - 1
    metrics = [
        ("events", f'=COUNTA(LOG!C{LOG_DATA_START}:C{last})'),
        ("evidence_files", f'=COUNTA(EVIDENCE!A{EVIDENCE_DATA_START}:A{ev_last})'),
        ("owned", f'=COUNTIF(LOG!G{LOG_DATA_START}:G{last},"owned")'),
        ("unowned", f'=COUNTIF(LOG!G{LOG_DATA_START}:G{last},"UNOWNED")'),
        ("renamed", f'=COUNTIF(LOG!I{LOG_DATA_START}:I{last},"RENAMED")'),
        ("chain_ok", f'=COUNTIF(LOG!U{LOG_DATA_START}:U{last},"OK")'),
        ("genesis", f'=COUNTIF(LOG!U{LOG_DATA_START}:U{last},"GENESIS")'),
        ("breaks", f'=COUNTIF(LOG!U{LOG_DATA_START}:U{last},"BREAK")+COUNTIF(LOG!U{LOG_DATA_START}:U{last},"BROKEN")'),
        ("blank_blame", f'=COUNTIFS(LOG!C{LOG_DATA_START}:C{last},"<>",LOG!E{LOG_DATA_START}:E{last},"")'),
        ("blank_long_outcome", f'=COUNTIFS(LOG!C{LOG_DATA_START}:C{last},"<>",LOG!K{LOG_DATA_START}:K{last},"")'),
    ]
    ws["A4"] = "metric"
    ws["B4"] = "count"
    _style_header(ws, 4, 2)
    for i, (name, formula) in enumerate(metrics, start=5):
        ws.cell(i, 1, name)
        ws.cell(i, 2, formula)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 12
    chart = BarChart()
    chart.type = "col"
    chart.title = "EmployeeLock DASH"
    chart.y_axis.title = "count"
    data = Reference(ws, min_col=2, min_row=4, max_row=14)
    cats = Reference(ws, min_col=1, min_row=5, max_row=14)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.legend = None
    ws.add_chart(chart, "D4")
    ws["A16"] = LIMITATION
    ws["A16"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A16:H20")
    _print_tabloid_landscape(ws)


def _build_lists(ws: Worksheet) -> None:
    ws["A1"] = "observer_hint"
    ws["B1"] = "exhibit_potential"
    ws["C1"] = "kind"
    _style_header(ws, 1, 3)
    for i, val in enumerate(("operator", "witness", "system", "third-party"), start=2):
        ws.cell(i, 1, val)
    for i, val in enumerate(("Y", "N", "maybe"), start=2):
        ws.cell(i, 2, val)
    for i, val in enumerate(("audio", "video", "photo", "document", "other", "none"), start=2):
        ws.cell(i, 3, val)
    for col, width in enumerate((16, 18, 14), start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _new_workbook() -> Workbook:
    wb = Workbook()
    cover = wb.active
    cover.title = "COVER"
    _build_cover(cover)
    log = wb.create_sheet("LOG")
    _build_log(log)
    evidence = wb.create_sheet("EVIDENCE")
    _build_evidence(evidence)
    chain = wb.create_sheet("CHAIN")
    _build_chain(chain)
    owners = wb.create_sheet("OWNERS")
    _build_owners(owners)
    dash = wb.create_sheet("DASH")
    _build_dash(dash)
    lists = wb.create_sheet("LISTS")
    _build_lists(lists)
    return wb


def init_workbook(path: str | Path, *, demo: bool = True, timestamp: str | None = None) -> Path:
    dest = Path(path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb = _new_workbook()
    wb.save(dest)
    if demo:
        ts1 = timestamp or utc_now()
        append_row(
            dest,
            event=DEMO_ROW_1["event"],
            result=DEMO_ROW_1["result"],
            blame=DEMO_ROW_1["blame"],
            owner=DEMO_ROW_1["owner"],
            short=DEMO_ROW_1["short"],
            long=DEMO_ROW_1["long"],
            renamed_from=DEMO_ROW_1["renamed_from"],
            confidence=DEMO_ROW_1["confidence"],
            observer=DEMO_ROW_1["observer"],
            exhibit=DEMO_ROW_1["exhibit"],
            timestamp=ts1,
        )
        ts2 = timestamp or utc_now()
        append_row(
            dest,
            event=DEMO_ROW_2["event"],
            result=DEMO_ROW_2["result"],
            blame=DEMO_ROW_2["blame"],
            owner=DEMO_ROW_2["owner"],
            short=DEMO_ROW_2["short"],
            long=DEMO_ROW_2["long"],
            renamed_from=DEMO_ROW_2["renamed_from"],
            confidence=DEMO_ROW_2["confidence"],
            observer=DEMO_ROW_2["observer"],
            exhibit=DEMO_ROW_2["exhibit"],
            timestamp=ts2,
        )
    return dest


def _next_log_row(ws: Worksheet) -> int:
    last = LOG_DATA_START + MAX_LOG - 1
    for row in range(LOG_DATA_START, last + 1):
        if not _cell_str(ws.cell(row, 3).value):
            return row
    raise ValueError(f"LOG full (MAX_LOG={MAX_LOG}); bump the builder")


def _next_evidence_row(ws: Worksheet) -> int:
    last = EVIDENCE_DATA_START + MAX_EVIDENCE - 1
    for row in range(EVIDENCE_DATA_START, last + 1):
        if not _cell_str(ws.cell(row, 1).value):
            return row
    raise ValueError(f"EVIDENCE full (MAX_EVIDENCE={MAX_EVIDENCE}); bump the builder")


def _entry_id_for_row(row: int) -> str:
    n = row - LOG_DATA_START + 1
    return f"EL-{n:04d}"


def _evidence_id_for_row(row: int) -> str:
    n = row - EVIDENCE_DATA_START + 1
    return f"EV-{n:04d}"


def _prev_hash_for_row(ws: Worksheet, row: int) -> str:
    if row == LOG_DATA_START:
        return GENESIS_PREV
    prev = _cell_str(ws.cell(row - 1, 20).value)
    return prev or GENESIS_PREV


def _refresh_owners(wb: Workbook) -> None:
    log = wb["LOG"]
    owners = wb["OWNERS"]
    last = LOG_DATA_START + MAX_LOG - 1
    seen: dict[str, list[int]] = {}
    blame_counts: dict[str, int] = {}
    for row in range(LOG_DATA_START, last + 1):
        event = _cell_str(log.cell(row, 3).value)
        if not event:
            continue
        owner = _cell_str(log.cell(row, 6).value)
        blame = _cell_str(log.cell(row, 5).value)
        if owner:
            seen.setdefault(owner, []).append(row)
        if blame:
            blame_counts[blame] = blame_counts.get(blame, 0) + 1
    for row in range(OWNERS_DATA_START, OWNERS_DATA_START + 64):
        owners.cell(row, 1).value = None
        owners.cell(row, 2).value = None
        owners.cell(row, 3).value = None
    for i, (name, rows) in enumerate(seen.items(), start=OWNERS_DATA_START):
        owners.cell(i, 1, name)
        owners.cell(i, 2, len(rows))
        owners.cell(i, 3, blame_counts.get(name, 0))


def _write_log_values(
    ws: Worksheet,
    row: int,
    *,
    entry_id: str,
    timestamp: str,
    event: str,
    result: str,
    blame: str,
    owner: str,
    renamed_from: str,
    short: str,
    long: str,
    cites: str,
    confidence: float,
    observer: str,
    exhibit: str,
    evidence_ids: str,
    file_sha256s: str,
    canonical: str,
    prev: str,
    digest: str,
) -> None:
    values = {
        1: entry_id,
        2: timestamp,
        3: event,
        4: result,
        5: blame,
        6: owner,
        8: renamed_from,
        10: short,
        11: long,
        12: cites,
        13: float(format_confidence(confidence)),
        14: observer,
        15: exhibit,
        16: evidence_ids,
        17: file_sha256s,
        18: canonical,
        19: prev,
        20: digest,
    }
    for col, value in values.items():
        cell = ws.cell(row, col, value)
        if col in INPUT_COLS:
            cell.font = FONT_BLUE
        else:
            cell.font = FONT_BODY
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row, 13).number_format = "0.000000"


def append_row(
    path: str | Path,
    *,
    event: str,
    result: str,
    blame: str = "",
    owner: str = "",
    short: str = "",
    long: str = "",
    renamed_from: str = "",
    cites: str = "",
    confidence: float = DEFAULT_CONFIDENCE,
    observer: str = DEFAULT_OBSERVER,
    exhibit: str = DEFAULT_EXHIBIT,
    timestamp: str | None = None,
    evidence_ids: str = "",
    file_sha256s: str = "",
) -> dict[str, Any]:
    dest = Path(path)
    wb = load_workbook(dest)
    log = wb["LOG"]
    row = _next_log_row(log)
    entry_id = _entry_id_for_row(row)
    ts = timestamp or utc_now()
    conf = float(confidence)
    if conf < 0.0 or conf > 1.0:
        raise ValueError("confidence must be in [0.0, 1.0]")
    prev = _prev_hash_for_row(log, row)
    fields = {
        "entry_id": entry_id,
        "timestamp": ts,
        "event": event,
        "result": result,
        "blame_placed": blame,
        "owner_named": owner,
        "renamed_from": renamed_from,
        "outcome_short": short,
        "outcome_long": long,
        "evidence_ids": evidence_ids,
        "confidence": conf,
        "observer": observer,
        "file_sha256s": file_sha256s,
        "prev_hash": prev,
    }
    canon, digest = row_hash(fields)
    _write_log_values(
        log,
        row,
        entry_id=entry_id,
        timestamp=ts,
        event=event,
        result=result,
        blame=blame,
        owner=owner,
        renamed_from=renamed_from,
        short=short,
        long=long,
        cites=cites,
        confidence=conf,
        observer=observer,
        exhibit=exhibit,
        evidence_ids=evidence_ids,
        file_sha256s=file_sha256s,
        canonical=canon,
        prev=prev,
        digest=digest,
    )
    _refresh_owners(wb)
    wb.save(dest)
    return {
        "entry_id": entry_id,
        "row": row,
        "row_hash": digest,
        "prev_hash": prev,
        "canonical": canon,
        "unowned": "UNOWNED" if not str(owner).strip() else "owned",
        "name_moved": "RENAMED" if str(renamed_from).strip() else "stable",
        "timestamp": ts,
    }


def import_files(
    path: str | Path,
    files: Sequence[str | Path],
    *,
    event: str,
    result: str,
    blame: str = "",
    owner: str = "",
    short: str = "",
    long: str = "",
    renamed_from: str = "",
    cites: str = "",
    confidence: float = DEFAULT_CONFIDENCE,
    observer: str = DEFAULT_OBSERVER,
    exhibit: str = DEFAULT_EXHIBIT,
    timestamp: str | None = None,
    kind: str | None = None,
    note: str = "",
) -> dict[str, Any]:
    dest = Path(path)
    wb = load_workbook(dest)
    log = wb["LOG"]
    ev = wb["EVIDENCE"]
    row = _next_log_row(log)
    entry_id = _entry_id_for_row(row)
    ts = timestamp or utc_now()
    ids: list[str] = []
    hashes: list[str] = []
    evidence_rows: list[dict[str, Any]] = []
    for file_path in files:
        src = Path(file_path)
        if not src.is_file():
            raise FileNotFoundError(str(src))
        digest, size = sha256_file(src)
        ev_row = _next_evidence_row(ev)
        ev_id = _evidence_id_for_row(ev_row)
        guessed = guess_kind(src, kind)
        stored = str(src.resolve())
        ev.cell(ev_row, 1, ev_id)
        ev.cell(ev_row, 2, ts)
        ev.cell(ev_row, 3, guessed)
        ev.cell(ev_row, 4, src.name)
        ev.cell(ev_row, 5, stored)
        ev.cell(ev_row, 6, size)
        ev.cell(ev_row, 7, digest)
        ev.cell(ev_row, 8, entry_id)
        ev.cell(ev_row, 9, note)
        ids.append(ev_id)
        hashes.append(digest)
        evidence_rows.append(
            {
                "evidence_id": ev_id,
                "kind": guessed,
                "original_name": src.name,
                "stored_path": stored,
                "size_bytes": size,
                "file_sha256": digest,
            }
        )
    evidence_ids = ",".join(ids)
    file_sha256s = ",".join(hashes)
    wb.save(dest)
    result_row = append_row(
        dest,
        event=event,
        result=result,
        blame=blame,
        owner=owner,
        short=short,
        long=long,
        renamed_from=renamed_from,
        cites=cites,
        confidence=confidence,
        observer=observer,
        exhibit=exhibit,
        timestamp=ts,
        evidence_ids=evidence_ids,
        file_sha256s=file_sha256s,
    )
    result_row["evidence"] = evidence_rows
    return result_row


def _iter_log_rows(ws: Worksheet) -> Iterable[int]:
    last = LOG_DATA_START + MAX_LOG - 1
    for row in range(LOG_DATA_START, last + 1):
        if _cell_str(ws.cell(row, 3).value):
            yield row


def read_log_row(ws: Worksheet, row: int) -> dict[str, Any]:
    conf_raw = ws.cell(row, 13).value
    try:
        conf = float(conf_raw) if conf_raw not in (None, "") else 0.0
    except (TypeError, ValueError):
        conf = 0.0
    return {
        "row": row,
        "entry_id": _cell_str(ws.cell(row, 1).value),
        "timestamp": _cell_str(ws.cell(row, 2).value),
        "event": _as_str(ws.cell(row, 3).value),
        "result": _as_str(ws.cell(row, 4).value),
        "blame_placed": _as_str(ws.cell(row, 5).value),
        "owner_named": _as_str(ws.cell(row, 6).value),
        "unowned_formula": _as_str(ws.cell(row, 7).value),
        "renamed_from": _as_str(ws.cell(row, 8).value),
        "name_moved_formula": _as_str(ws.cell(row, 9).value),
        "outcome_short": _as_str(ws.cell(row, 10).value),
        "outcome_long": _as_str(ws.cell(row, 11).value),
        "cites": _as_str(ws.cell(row, 12).value),
        "confidence": conf,
        "observer": _as_str(ws.cell(row, 14).value),
        "exhibit_potential": _as_str(ws.cell(row, 15).value),
        "evidence_ids": _as_str(ws.cell(row, 16).value),
        "file_sha256s": _as_str(ws.cell(row, 17).value),
        "canonical": _as_str(ws.cell(row, 18).value),
        "prev_hash": _cell_str(ws.cell(row, 19).value),
        "row_hash": _cell_str(ws.cell(row, 20).value),
        "link_ok_formula": _as_str(ws.cell(row, 21).value),
    }


def hashed_payload(row: dict[str, Any]) -> dict[str, Any]:
    return {k: row[k] if k != "timestamp" else row["timestamp"] for k in HASHED_FIELDS}


def verify_workbook(path: str | Path) -> dict[str, Any]:
    dest = Path(path)
    wb = load_workbook(dest, data_only=False)
    log = wb["LOG"]
    errors: list[str] = []
    missing_files: list[str] = []
    unowned = 0
    rows_out: list[dict[str, Any]] = []
    prev_stored = GENESIS_PREV
    for i, row in enumerate(_iter_log_rows(log)):
        rec = read_log_row(log, row)
        fields = {
            "entry_id": rec["entry_id"],
            "timestamp": rec["timestamp"],
            "event": rec["event"],
            "result": rec["result"],
            "blame_placed": rec["blame_placed"],
            "owner_named": rec["owner_named"],
            "renamed_from": rec["renamed_from"],
            "outcome_short": rec["outcome_short"],
            "outcome_long": rec["outcome_long"],
            "evidence_ids": rec["evidence_ids"],
            "confidence": rec["confidence"],
            "observer": rec["observer"],
            "file_sha256s": rec["file_sha256s"],
            "prev_hash": rec["prev_hash"],
        }
        canon, digest = row_hash(fields)
        if rec["row_hash"] != digest:
            errors.append(
                f"{rec['entry_id']}: stored row_hash {rec['row_hash']} != recomputed {digest}"
            )
        expected_prev = GENESIS_PREV if i == 0 else prev_stored
        if rec["prev_hash"] != expected_prev:
            errors.append(
                f"{rec['entry_id']}: prev_hash {rec['prev_hash']} != expected {expected_prev}"
            )
        if not str(rec["owner_named"]).strip():
            unowned += 1
        rec["recomputed"] = digest
        rec["canonical_recomputed"] = canon
        rec["unowned"] = "UNOWNED" if not str(rec["owner_named"]).strip() else "owned"
        rec["name_moved"] = "RENAMED" if str(rec["renamed_from"]).strip() else "stable"
        rows_out.append(rec)
        prev_stored = rec["row_hash"] or digest

    if "EVIDENCE" in wb.sheetnames:
        ev = wb["EVIDENCE"]
        last = EVIDENCE_DATA_START + MAX_EVIDENCE - 1
        for row in range(EVIDENCE_DATA_START, last + 1):
            ev_id = _cell_str(ev.cell(row, 1).value)
            if not ev_id:
                continue
            stored_path = _cell_str(ev.cell(row, 5).value)
            stored_hash = _cell_str(ev.cell(row, 7).value)
            if not stored_path or not os.path.isfile(stored_path):
                missing_files.append(stored_path or f"{ev_id}: (empty path)")
                continue
            digest, _size = sha256_file(Path(stored_path))
            if stored_hash and digest != stored_hash:
                errors.append(f"{ev_id}: file bytes changed at {stored_path}")

    ok = True
    for err in errors:
        if "row_hash" in err or "prev_hash" in err:
            ok = False
            break
    result = {
        "ok": ok,
        "rows": len(rows_out),
        "errors": errors,
        "missing_files": missing_files,
        "unowned": unowned,
        "spec": SPEC_STRING,
        "version": ENGINE_VERSION,
        "limitation": LIMITATION,
    }
    return result


def list_rows(path: str | Path) -> list[dict[str, Any]]:
    wb = load_workbook(Path(path), data_only=False)
    log = wb["LOG"]
    out = []
    for row in _iter_log_rows(log):
        rec = read_log_row(log, row)
        rec["unowned"] = "UNOWNED" if not str(rec["owner_named"]).strip() else "owned"
        rec["name_moved"] = "RENAMED" if str(rec["renamed_from"]).strip() else "stable"
        out.append(rec)
    return out


def dash_counts(path: str | Path) -> dict[str, int]:
    rows = list_rows(path)
    report = verify_workbook(path)
    owned = sum(1 for r in rows if r["unowned"] == "owned")
    renamed = sum(1 for r in rows if r["name_moved"] == "RENAMED")
    genesis = 1 if rows else 0
    chain_ok = max(0, len(rows) - genesis) if report["ok"] else 0
    blank_blame = sum(1 for r in rows if not str(r["blame_placed"]).strip())
    blank_long = sum(1 for r in rows if not str(r["outcome_long"]).strip())
    return {
        "events": len(rows),
        "evidence_files": _count_evidence(path),
        "owned": owned,
        "unowned": report["unowned"],
        "renamed": renamed,
        "chain_ok": chain_ok,
        "genesis": genesis,
        "breaks": 0 if report["ok"] else 1,
        "blank_blame": blank_blame,
        "blank_long_outcome": blank_long,
        "missing_files": len(report["missing_files"]),
    }


def _count_evidence(path: str | Path) -> int:
    wb = load_workbook(Path(path), data_only=False)
    if "EVIDENCE" not in wb.sheetnames:
        return 0
    ev = wb["EVIDENCE"]
    n = 0
    last = EVIDENCE_DATA_START + MAX_EVIDENCE - 1
    for row in range(EVIDENCE_DATA_START, last + 1):
        if _cell_str(ev.cell(row, 1).value):
            n += 1
    return n
